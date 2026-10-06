#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Encontros do Ministro de Estado (ME) e do Secretario-Geral (SG) do MRE com a
Fazenda / Economia, o Planejamento e Orcamento e a Casa Civil, nas agendas
publicadas em gov.br (nov/2020 a mai/2023), com triagem analoga a do relatorio
"MRE_encontros_restritos_Fazenda_MPO_CasaCivil" (fonte e-Agendas, 2023-2026).

Saidas: agendas_mre_encontros.xlsx (lista completa + triagem) e
        agendas_mre_encontros.json (linhas mantidas, para o relatorio .docx).
"""
import csv, json, re, sys, unicodedata, datetime as dt
from collections import OrderedDict

CSV_IN = "agendas_mre.csv"
XLSX = "agendas_mre_encontros.xlsx"
JSON_OUT = "agendas_mre_encontros.json"

def norm(s):
    s = unicodedata.normalize("NFKD", s or "")
    return re.sub(r"\s+", " ", "".join(c for c in s if not unicodedata.combining(c)).lower()).strip()

AUT = {  # autoridade no CSV -> (nome, agenda, cargo curto)
    "Ernesto Araujo": ("Ernesto Fraga Araújo", "Ministro de Estado", "ME"),
    "Carlos Alberto Franco Franca": ("Carlos Alberto Franco França", "Ministro de Estado", "ME"),
    "Ministro das Relacoes Exteriores (atual)": ("Mauro Vieira", "Ministro de Estado", "ME"),
    "Otavio Brandelli": ("Otávio Brandelli", "Secretário-Geral", "SG"),
    "Fernando Simas Magalhaes": ("Fernando Simas Magalhães", "Secretário-Geral", "SG"),
}

# 1) Selecao de candidatos: termos do filtro pedido + varredura de orgaos/cargos
CAND = re.compile(
    r"\beconomia\b|\bfazenda\b|\bplanejamento\b|\bcasa civil\b|\bguedes\b|\bciro nogueira\b|\bbraga netto\b|\bramos\b|"
    r"\borcamentos?\b|\borcamentari[oa]s?\b|tesouro|receita federal|secretaria de orcamento|\bsof\b|colnago|funchal|"
    r"guaranys|waldery|montezano|carlos da costa|junta de execucao|\bjeo\b|secretaria de governo|\bsegov\b|celio faria|"
    r"jonathas|belchior|secretaria-geral da presidencia|secretaria geral/pr|\bsgpr\b|advocacia-geral|controladoria|"
    r"gestao e da inovacao|\bmgi\b|haddad|tebet|durigan|rosito|rui costa|mercadante|dweck")

# 2) Regras de triagem, em ordem. (regex sobre descricao normalizada, classificacao, motivo)
RULES = [
    (r"marcelo ramos|faro ramos|helio vitor ramos|ramos araujo|ramos horta|edite ramos|planejamento diplomatico|"
     r"fazenda da esperanca|casa civil de mocambique|economia e saude global|sergio rodrigues dos santos",
     "Excluído — falso positivo lexical", "Homônimo ou termo fora de contexto (Ramos; 'Planejamento Diplomático'; 'Fazenda da Esperança'; G20 'Economia e Saúde Global')."),
    (r"^posse\b|cerimonia de posse|^cerimonia\b|solenidade|sancao da lei",
     "Excluído — porte (cerimônia pública)", "Cerimônia ou posse: ato público amplo, análogo ao critério de porte do relatório."),
    (r"eslovenia|\bg20\b|\bbrics\b|cupula|itaipu|\bcop ?\d|mocambique|sao tome|embaixador d[aoe] |timor",
     "Excluído — tema de política externa", "Título indica tema de política externa (bilateral, cúpula, interlocutor estrangeiro)."),
    (r"secretaria de governo|\bsegov\b|secretaria-geral da presidencia|secretaria geral/pr|advocacia-geral|controladoria|"
     r"bndes|chefe do gabinete da presidencia|gestao e da inovacao|\bmgi\b|secretaria-geral da presidencia",
     "Excluído — órgão fora do escopo", "Órgão não coberto pelo relatório (SEGOV, SGPR, AGU, CGU, BNDES, MGI, Gabinete da PR)."),
]

def orgao(nd):
    if re.search(r"casa civil|braga netto|ciro nogueira|rui costa|jonathas|belchior|sergio pereira|thiago meirelles|bruno (cesar )?grossi|hott junior", nd):
        return "Casa Civil"
    if re.search(r"receita federal", nd): return "Ministério da Economia — Receita Federal"
    if re.search(r"orcamento federal|culau", nd): return "Ministério da Economia — Secretaria de Orçamento Federal"
    if re.search(r"economia|guedes|guaranys|carlos da costa|daniella marques|caio mario|sultani|fendt", nd):
        return "Ministério da Economia"
    if re.search(r"fazenda|haddad", nd): return "Ministério da Fazenda"
    if re.search(r"planejamento e orcamento|tebet", nd): return "Ministério do Planejamento e Orçamento"
    if re.search(r"ramos", nd): return "Casa Civil (Luiz Eduardo Ramos)"
    return "(outro)"

def observacao(r, nd, cls):
    obs = []
    if nd.startswith("acompanha o ministro") or nd.startswith("acompanha o senhor ministro"):
        obs.append("Registro da Secretaria-Geral do mesmo evento da agenda do Ministro de Estado.")
    if "telefonema" in nd: obs.append("Telefonema.")
    if "videoconferencia" in nd or "videoconf" in norm(r["local"]): obs.append("Videoconferência.")
    if "almoco" in nd: obs.append("Almoço; porte não verificável (a agenda gov.br não lista participantes).")
    if "comercio exterior e assuntos internacionais" in nd:
        obs.append("Mantido por analogia à decisão do item 1.3 do relatório: o termo 'Internacionais' integra o cargo do interlocutor.")
    if "receita federal" in nd:
        obs.append("A Receita Federal integra o Ministério da Economia (sucessor da Fazenda); relevância orçamentária indireta.")
    if "enap" in norm(r["local"]): obs.append("Local (ENAP) sugere evento; porte não verificável.")
    if re.search(r"luiz eduardo ramos", nd) and "casa civil" in nd:
        obs.append("Ramos chefiou a Casa Civil de 29/03/2021 a 03/08/2021.")
    if "reuniao de coordenacao" in nd: obs.append("Reunião de coordenação sem pauta declarada.")
    if "presidencia do senado" in nd: obs.append("Também presente diretor da Presidência do Senado Federal.")
    if r["data"] == "2021-09-09" and "caio mario" in nd: obs.append("Reunião e almoço do mesmo dia: dois registros, uma visita.")
    if r["autoridade"] == "Ministro das Relacoes Exteriores (atual)" and cls == "Mantido":
        obs.append("Período anterior a 08/05/2023, não coberto pela agenda do ME no e-Agendas.")
    return " ".join(obs)

def main():
    rows = list(csv.DictReader(open(CSV_IN, encoding="utf-8")))
    cands = []
    for r in rows:
        nd = norm(r["descricao"])
        if not CAND.search(nd): continue
        cls, motivo = "Mantido", "Interlocutor de órgão do escopo; título sem tema de política externa; composição aparentemente restrita."
        for rx, c, m in RULES:
            if re.search(rx, nd): cls, motivo = c, m; break
        if cls == "Mantido" and orgao(nd) == "(outro)":
            cls, motivo = "Excluído — órgão fora do escopo", "Nenhum órgão do escopo identificado no título."
        nome, agenda, cargo = AUT[r["autoridade"]]
        cands.append(OrderedDict([
            ("data", r["data"]), ("hora", r["hora_inicio"]), ("agenda", agenda), ("cargo", cargo), ("autoridade", nome),
            ("orgao", orgao(nd) if cls == "Mantido" else ("Casa Civil (Luiz Eduardo Ramos, SEGOV/SGPR)" if "ramos" in nd and "secretaria" in nd else orgao(nd))),
            ("titulo", r["descricao"].strip()), ("local", r["local"].strip()),
            ("classificacao", cls), ("motivo", motivo), ("observacao", observacao(r, nd, cls)),
            ("url", r["url"]),
        ]))
    cands.sort(key=lambda c: (c["data"], c["hora"], c["cargo"]))
    # eventos: ME+SG no mesmo dia/hora = 1 evento
    ev = {}; n = 0
    for c in cands:
        if c["classificacao"] != "Mantido": c["evento"] = ""; continue
        k = (c["data"], c["hora"])
        if k not in ev: n += 1; ev[k] = n
        c["evento"] = ev[k]
    for i, c in enumerate(cands, 1): c["n"] = i
    kept = [c for c in cands if c["classificacao"] == "Mantido"]
    json.dump(kept, open(JSON_OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    write_xlsx(cands, kept, n)
    from collections import Counter
    print("candidatos: %d | mantidos: %d (%d eventos distintos)" % (len(cands), len(kept), n))
    for k, v in Counter(c["classificacao"] for c in cands).most_common(): print("  %-45s %3d" % (k, v))
    print("  mantidos por agenda:", dict(Counter(c["cargo"] for c in kept)))
    print("  mantidos por autoridade:", dict(Counter(c["autoridade"] for c in kept)))
    if "--print" in sys.argv:
        for c in cands:
            print("%3d %s %s %-2s %-45s | %s | %s" % (c["n"], c["data"], c["hora"], c["cargo"], c["classificacao"][:45], c["titulo"][:80], c["observacao"][:60]))

def write_xlsx(cands, kept, nev):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    F = "Arial"; hf = Font(name=F, bold=True, color="FFFFFF", size=10); hfill = PatternFill("solid", fgColor="1F3864")
    body = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
    def header(ws, cols, widths):
        ws.append(cols)
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i); cell.font = hf; cell.fill = hfill; cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = widths[i-1]
        ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 30
    def finish(ws, n, ncols):
        for r in range(2, n + 2):
            for c in range(1, ncols + 1):
                cell = ws.cell(row=r, column=c); cell.font = body; cell.alignment = Alignment(wrap_text=True, vertical="top")
                if c == 2: cell.number_format = "DD/MM/YYYY"
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(ncols), n + 1)
    cols = ["Nº", "Data", "Hora", "Agenda", "Autoridade", "Órgão identificado", "Compromisso (título em gov.br)", "Local",
            "Classificação", "Motivo", "Observações", "Evento", "URL (gov.br)"]
    widths = [5, 11, 7, 17, 26, 30, 70, 32, 34, 48, 48, 8, 60]
    wb = Workbook(); ws = wb.active; ws.title = "Encontros"
    header(ws, cols, widths)
    for c in cands:
        ws.append([c["n"], dt.date.fromisoformat(c["data"]), c["hora"], c["agenda"], c["autoridade"], c["orgao"], c["titulo"], c["local"],
                   c["classificacao"], c["motivo"], c["observacao"], c["evento"], c["url"]])
    finish(ws, len(cands), len(cols)); N = len(cands) + 1
    wm = wb.create_sheet("Mantidos"); header(wm, cols, widths)
    for c in kept:
        wm.append([c["n"], dt.date.fromisoformat(c["data"]), c["hora"], c["agenda"], c["autoridade"], c["orgao"], c["titulo"], c["local"],
                   c["classificacao"], c["motivo"], c["observacao"], c["evento"], c["url"]])
    finish(wm, len(kept), len(cols))
    wr = wb.create_sheet("Resumo", 0)
    wr.column_dimensions["A"].width = 46
    for col in "BCD": wr.column_dimensions[col].width = 16
    wr["A1"] = "Encontros do ME e do SG do MRE com Fazenda/Economia, Planejamento e Orçamento e Casa Civil — agendas gov.br"; wr["A1"].font = Font(name=F, bold=True, size=12)
    wr["A2"] = "Período: 01/11/2020 a 31/12/2022 (Araújo, França, Brandelli, Simas Magalhães) e 01/01/2023 a 07/05/2023 (Mauro Vieira). Extraído em 06/10/2026."; wr["A2"].font = Font(name=F, size=9, italic=True)
    for i, h in enumerate(["Classificação", "Ministro de Estado", "Secretário-Geral", "Total"], 1):
        c = wr.cell(row=4, column=i, value=h); c.font = hf; c.fill = hfill
    classes = ["Mantido", "Excluído — falso positivo lexical", "Excluído — porte (cerimônia pública)",
               "Excluído — tema de política externa", "Excluído — órgão fora do escopo"]
    for i, cl in enumerate(classes):
        r = 5 + i
        wr.cell(row=r, column=1, value=cl).font = body
        wr.cell(row=r, column=2, value='=COUNTIFS(Encontros!$I$2:$I$%d,$A%d,Encontros!$D$2:$D$%d,"Ministro de Estado")' % (N, r, N)).font = body
        wr.cell(row=r, column=3, value='=COUNTIFS(Encontros!$I$2:$I$%d,$A%d,Encontros!$D$2:$D$%d,"Secretário-Geral")' % (N, r, N)).font = body
        wr.cell(row=r, column=4, value="=B%d+C%d" % (r, r)).font = body
    rt = 5 + len(classes)
    wr.cell(row=rt, column=1, value="Total de candidatos").font = bold
    for c in range(2, 5):
        L = get_column_letter(c); wr.cell(row=rt, column=c, value="=SUM(%s5:%s%d)" % (L, L, rt - 1)).font = bold
    wr.cell(row=rt + 2, column=1, value="Eventos distintos mantidos (ME+SG no mesmo horário = 1)").font = body
    wr.cell(row=rt + 2, column=2, value=nev).font = body
    wr.cell(row=rt + 2, column=3, value="valor calculado pelo script (coluna Evento da aba Encontros)").font = Font(name=F, size=9, italic=True)
    notas = ["", "Como ler", "• 'Candidatos' = compromissos cujo título cita um dos órgãos, seus titulares ou cargos (lista de termos na aba Método).",
             "• A triagem replica, no que a fonte permite, os critérios do relatório e-Agendas: tema de política externa, porte, órgão do escopo.",
             "• A agenda gov.br não lista participantes: o critério de porte só pôde ser aplicado a cerimônias/posses. Almoços foram mantidos, como no relatório.",
             "• Mapeamento de órgãos 2019–2022: Ministério da Economia = sucessor de Fazenda e Planejamento/Orçamento; Receita Federal e SOF integram a Economia.",
             "• Fora do escopo (como no relatório): Secretaria de Governo, Secretaria-Geral da Presidência, AGU, CGU, BNDES, MGI."]
    for i, t in enumerate(notas):
        wr.cell(row=rt + 4 + i, column=1, value=t).font = bold if t == "Como ler" else Font(name=F, size=9)
    wmet = wb.create_sheet("Método"); wmet.column_dimensions["A"].width = 120
    met = [("Método", True),
           ("Fonte: agendas diárias públicas em www.gov.br/mre (agendas de autoridades anteriores e agenda do ministro), uma página por dia; 3.291 dias requisitados, 4.614 compromissos extraídos (agendas_mre.csv).", False),
           ("Seleção de candidatos (regex sobre o título, sem acentos): economia, fazenda, planejamento, casa civil, guedes, ciro nogueira, braga netto, ramos, orçamento(s), orçamentário(a); mais: tesouro, receita federal, secretaria de orçamento, SOF, colnago, funchal, guaranys, waldery, montezano, carlos da costa, JEO, secretaria de governo, SEGOV, celio faria, jonathas, belchior, secretaria-geral da presidência, SGPR, advocacia-geral, controladoria, MGI, haddad, tebet, durigan, rosito, rui costa, mercadante, dweck.", False),
           ("", False), ("Triagem, em ordem:", True),
           ("1. Falso positivo lexical — homônimos de 'Ramos' (dep. Marcelo Ramos; emb. Luís Faro Ramos; emb. Hélio Vitor Ramos Filho; min. Marcelo Ramos Araújo; José Ramos-Horta; Edite Ramos), 'Planejamento Diplomático' (assessoria interna do MRE), 'Fazenda da Esperança', sessão do G20 'Economia e Saúde Global'.", False),
           ("2. Porte — posses e cerimônias públicas.", False),
           ("3. Tema de política externa — bilateral, cúpulas, interlocutor estrangeiro (ex.: reunião da Casa Civil com o embaixador da Eslovênia; Casa Civil de Moçambique).", False),
           ("4. Órgão fora do escopo — Secretaria de Governo (Ramos, mar/2021; Flávia Arruda), Secretaria-Geral da Presidência (Ramos, a partir de 03/08/2021; Onyx Lorenzoni; Jorge Oliveira), AGU, CGU, BNDES, MGI, Gabinete da PR.", False),
           ("5. Mantido — o restante.", False),
           ("", False), ("Datas de referência (verificadas em 06/10/2026): Braga Netto na Casa Civil de 18/02/2020 a 29/03/2021; Luiz Eduardo Ramos na Casa Civil de 29/03/2021 a 03/08/2021 e na Secretaria-Geral da PR a partir de 03/08/2021; Ciro Nogueira na Casa Civil a partir de 04/08/2021.", False),
           ("", False), ("Limitações: o título é a única informação disponível (sem participantes nem assunto); o porte real das reuniões é desconhecido; a agenda publicada é curada e não registra despachos internos.", False),
           ("Reprodutibilidade: scripts/mre_encontros.py (gera este arquivo a partir de agendas_mre.csv).", False)]
    for i, (t, b) in enumerate(met, 1):
        c = wmet.cell(row=i, column=1, value=t); c.font = Font(name=F, size=10, bold=b); c.alignment = Alignment(wrap_text=True, vertical="top")
    # LibreOffice nao esta disponivel para gravar valores em cache: forca o
    # recalculo completo ao abrir (Excel / Google Sheets / LibreOffice).
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(XLSX)

if __name__ == "__main__":
    main()
